"""Regression tests for the second API review: upload guards, export notes and headers, error
mapping, outcome-leakage checks, request limits, and interval and validation budgets."""

from __future__ import annotations

import io
import zipfile

import numpy as np
import pytest
from fastapi.testclient import TestClient

import survival_toolkit.app as app_module
from survival_toolkit.app import app

client = TestClient(app, base_url="http://127.0.0.1")


def _detail(response) -> str:
    detail = response.json().get("detail")
    if isinstance(detail, list):
        return " | ".join(str(item.get("msg")) for item in detail)
    return str(detail)


def _must_not_load(*args, **kwargs):
    raise AssertionError("the upload must be refused before the table is parsed")


# ── Upload guards: workbook parts, Parquet text types, legacy Excel content ──

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_SHARED_STRINGS_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"


def _workbook(
    *,
    shared_name: str = "xl/sharedStrings.xml",
    sheet_name: str = "xl/worksheets/sheet1.xml",
    n_rows: int = 300,
    note_length: int = 5000,
    shared_prolog: str = "",
    shared_encoding: str = "utf-8",
    extra_members: int = 0,
) -> bytes:
    """A workbook whose note cells all use one shared string, with the parts stored under the given names."""
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        f'<Override PartName="/{sheet_name}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        f'<Override PartName="/{shared_name}" ContentType="{_SHARED_STRINGS_TYPE}"/>'
        "</Types>"
    )
    package_rels = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{_PACKAGE_REL_NS}">'
        f'<Relationship Id="rId1" Type="{_REL_NS}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
    )
    workbook = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="{_MAIN_NS}" xmlns:r="{_REL_NS}">'
        '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    workbook_rels = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{_PACKAGE_REL_NS}">'
        f'<Relationship Id="rId1" Type="{_REL_NS}/worksheet" Target="/{sheet_name}"/>'
        f'<Relationship Id="rId2" Type="{_REL_NS}/sharedStrings" Target="/{shared_name}"/></Relationships>'
    )
    declared = "UTF-16" if shared_encoding == "utf-16" else "UTF-8"
    shared = (
        f'<?xml version="1.0" encoding="{declared}" standalone="yes"?>{shared_prolog}<sst xmlns="{_MAIN_NS}" count="4" uniqueCount="4">'
        f"<si><t>{'x' * note_length}</t></si><si><t>os_months</t></si><si><t>os_event</t></si><si><t>note</t></si></sst>"
    ).encode(shared_encoding)
    rows = ['<row r="1"><c r="A1" t="s"><v>1</v></c><c r="B1" t="s"><v>2</v></c><c r="C1" t="s"><v>3</v></c></row>']
    for index in range(2, n_rows + 2):
        rows.append(
            f'<row r="{index}"><c r="A{index}"><v>{index % 97 + 1}</v></c><c r="B{index}"><v>{index % 2}</v></c>'
            f'<c r="C{index}" t="s"><v>0</v></c></row>'
        )
    sheet = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="{_MAIN_NS}"><sheetData>{"".join(rows)}</sheetData></worksheet>'
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", package_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr(shared_name, shared)
        archive.writestr(sheet_name, sheet)
        for index in range(extra_members):
            archive.writestr(f"xl/media/image{index}.png", b"\x89PNG\r\n\x1a\n" + bytes(16))
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("shared_name", "sheet_name"),
    [("xl/strings.bin", "xl/worksheets/sheet1.xml"), ("xl/sharedStrings.xml", "xl/worksheets/sheet1.bin"), ("xl/s.dat", "xl/w.dat")],
    ids=["shared-strings", "sheet", "both"],
)
def test_workbook_parts_under_any_name_are_measured(monkeypatch: pytest.MonkeyPatch, shared_name: str, sheet_name: str) -> None:
    pytest.importorskip("openpyxl")
    payload = _workbook(shared_name=shared_name, sheet_name=sheet_name)
    # openpyxl finds the parts through the content types and relationships, whatever their names.
    accepted = client.post("/api/upload", files={"file": ("renamed.xlsx", payload, "application/octet-stream")})
    assert accepted.status_code == 200, accepted.text
    assert [column["name"] for column in accepted.json()["columns"]] == ["os_months", "os_event", "note"]

    monkeypatch.setattr(app_module, "_MAX_UPLOAD_TEXT_CHARS", 1_000_000)
    monkeypatch.setattr(app_module, "load_dataframe_from_path", _must_not_load)
    refused = client.post("/api/upload", files={"file": ("renamed.xlsx", payload, "application/octet-stream")})
    assert refused.status_code == 413, refused.text
    assert "of text" in _detail(refused)


@pytest.mark.parametrize(
    ("shared_name", "prolog", "encoding"),
    [
        ("xl/strings.bin", '<!DOCTYPE sst [<!ENTITY a "aaaa">]>', "utf-8"),
        ("xl/sharedStrings.xml", "<!--" + " padding " * 1000 + '--><!DOCTYPE sst [<!ENTITY a "aaaa">]>', "utf-8"),
        ("xl/sharedStrings.xml", '<!DOCTYPE sst [<!ENTITY a "aaaa">]>', "utf-16"),
    ],
    ids=["part-not-named-xml", "after-a-long-comment", "utf-16"],
)
def test_workbook_document_types_are_refused_in_every_xml_part(
    monkeypatch: pytest.MonkeyPatch, shared_name: str, prolog: str, encoding: str
) -> None:
    payload = _workbook(shared_name=shared_name, n_rows=20, note_length=10, shared_prolog=prolog, shared_encoding=encoding)
    monkeypatch.setattr(app_module, "load_dataframe_from_path", _must_not_load)
    response = client.post("/api/upload", files={"file": ("doctype.xlsx", payload, "application/octet-stream")})
    assert response.status_code == 400, response.text
    assert "not allowed" in _detail(response)


def test_workbooks_with_too_many_parts_are_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(app_module, "_MAX_XLSX_MEMBERS", 8)
    monkeypatch.setattr(app_module, "load_dataframe_from_path", _must_not_load)
    payload = _workbook(n_rows=5, note_length=5, extra_members=10)
    response = client.post("/api/upload", files={"file": ("many.xlsx", payload, "application/octet-stream")})
    assert response.status_code == 400, response.text
    assert "parts" in _detail(response)


def _parquet_text_table(kind: str, n_rows: int = 2000, width: int = 2000):
    pa = pytest.importorskip("pyarrow")
    if kind == "string_view":
        values = pa.array(["x" * width] * n_rows, type=pa.string_view())
    elif kind == "binary_view":
        values = pa.array([b"y" * width] * n_rows, type=pa.binary_view())
    elif kind == "fixed_size_binary":
        values = pa.array([b"z" * width] * n_rows, type=pa.binary(width))
    elif kind == "json":
        if not hasattr(pa, "json_"):
            pytest.skip("this pyarrow has no JSON extension type")
        values = pa.array(['"' + "j" * width + '"'] * n_rows, type=pa.json_())
    else:  # pragma: no cover - test parameter typo
        raise AssertionError(kind)
    return pa.table({"os_months": np.linspace(1.0, 50.0, n_rows), "os_event": np.arange(n_rows) % 2, "note": values})


@pytest.mark.parametrize(
    ("kind", "status", "message"),
    [
        ("string_view", 413, "of text"),
        ("binary_view", 413, "of text"),
        ("fixed_size_binary", 413, "of text"),
        # Extension types cannot be read with their dictionaries kept, so they are refused by name.
        ("json", 400, "not supported"),
    ],
)
def test_parquet_text_of_every_non_numeric_type_is_measured(
    monkeypatch: pytest.MonkeyPatch, kind: str, status: int, message: str
) -> None:
    pq = pytest.importorskip("pyarrow.parquet")
    buffer = io.BytesIO()
    pq.write_table(_parquet_text_table(kind), buffer, compression="zstd")
    assert len(buffer.getvalue()) < 200_000

    monkeypatch.setattr(app_module, "_MAX_UPLOAD_TEXT_CHARS", 1_000_000)
    monkeypatch.setattr(app_module, "load_dataframe_from_path", _must_not_load)
    response = client.post("/api/upload", files={"file": (f"{kind}.parquet", buffer.getvalue(), "application/octet-stream")})
    assert response.status_code == status, response.text
    assert message in _detail(response)


def test_parquet_text_measurement_counts_each_type_once_per_cell(tmp_path) -> None:
    pq = pytest.importorskip("pyarrow.parquet")
    for kind in ("string_view", "binary_view", "fixed_size_binary"):
        path = tmp_path / f"{kind}.parquet"
        pq.write_table(_parquet_text_table(kind, n_rows=50, width=100), path)
        assert app_module._parquet_text_bytes(path, 10**9) == 50 * 100, kind


_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def test_legacy_excel_content_is_routed_by_signature_whatever_its_suffix(monkeypatch: pytest.MonkeyPatch) -> None:
    import pandas as pd

    payload = _OLE_SIGNATURE + bytes(504)
    checked: list[str] = []

    def _huge_text(path, limit):
        checked.append(str(path))
        return limit + 1

    monkeypatch.setattr(app_module, "_xls_text_chars", _huge_text)
    monkeypatch.setattr(app_module, "load_dataframe_from_path", _must_not_load)
    refused = client.post("/api/upload", files={"file": ("legacy.xlsx", payload, "application/octet-stream")})
    assert refused.status_code == 413, refused.text
    assert checked, "the legacy workbook check must run for a legacy file named .xlsx"

    loaded = pd.DataFrame({"os_months": [1.0, 2.0, 3.0], "os_event": [1, 0, 1]})
    monkeypatch.setattr(app_module, "_xls_text_chars", lambda path, limit: 0)
    monkeypatch.setattr(app_module, "load_dataframe_from_path", lambda *args, **kwargs: loaded.copy())
    accepted = client.post("/api/upload", files={"file": ("legacy.xlsx", payload, "application/octet-stream")})
    assert accepted.status_code == 200, accepted.text


# ── Replay and provenance notes stay within the export note limit ──


def _wide_request_config(n_features: int = 250, *, name_length: int = 15) -> dict:
    return {
        "model_type": "compare",
        "time_column": "os_months",
        "event_column": "os_event",
        "event_positive_value": 1,
        "features": [f"ENSG{index:0{name_length - 4}d}" for index in range(n_features)],
        "categorical_features": [f"CAT{index:0{name_length - 3}d}" for index in range(n_features)],
        "evaluation_strategy": "holdout",
        "random_state": 42,
        "n_estimators": 100,
        "max_depth": None,
        "learning_rate": 0.1,
        "random_seed": 7,
        "epochs": 10,
        "batch_size": 32,
        "hidden_layers": [16],
        "dropout": 0.1,
        "early_stopping_patience": 5,
        "early_stopping_min_delta": 0.0,
    }


def test_replay_notes_of_wide_comparisons_fit_the_export_note_limit() -> None:
    config = _wide_request_config()
    for notes in (
        app_module._ml_replay_notes(config, dataset_filename="tcga_expression.csv"),
        app_module._dl_replay_notes(config, dataset_filename="tcga_expression.csv", resolved_analysis={"evaluation_note": "x" * 5000}),
    ):
        assert notes and all(len(note) <= 4000 for note in notes), [len(note) for note in notes]
        features_note = next(note for note in notes if note.startswith("Replay features:"))
        assert "ENSG00000000000" in features_note and "250 in total" in features_note
        response = client.post(
            "/api/export-table",
            json={"rows": [{"Model": "RSF", "C-index": 0.71}], "format": "csv", "style": "journal", "notes": notes},
        )
        assert response.status_code == 200, response.text


def test_short_replay_feature_lists_are_written_in_full() -> None:
    notes = app_module._ml_replay_notes(
        {"features": ["age", "stage"], "categorical_features": ["stage"]}, dataset_filename="cohort.csv"
    )
    assert "Replay features: age, stage." in notes
    assert "Replay categorical features: stage." in notes


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "markdown", "latex", "docx"])
def test_provenance_notes_of_wide_configs_are_summarised(fmt: str) -> None:
    config = _wide_request_config(1000, name_length=30)
    notes = app_module._export_provenance_notes(
        {"dataset_hash": "abc123", "request_config": config, "analysis": {"evaluation_mode": "holdout", "folds": list(range(500))}}
    )
    assert all(len(note) <= 4000 for note in notes), [len(note) for note in notes]
    replay = next(note for note in notes if note.startswith("Replay request_config:"))
    assert "1,000 in total" in replay and '"random_state": 42' in replay
    response = client.post(
        "/api/export-table",
        json={
            "rows": [{"Model": "RSF", "C-index": 0.71}],
            "format": fmt,
            "style": "journal",
            "provenance": {"dataset_hash": "abc123", "request_config": config},
        },
    )
    assert response.status_code == 200, response.text
    if fmt == "xlsx":
        from openpyxl import load_workbook

        worksheet = load_workbook(io.BytesIO(response.content)).active
        assert all(len(str(cell.value or "")) <= 4000 for row in worksheet.iter_rows() for cell in row)


def test_small_provenance_configs_are_recorded_verbatim() -> None:
    notes = app_module._export_provenance_notes({"request_config": {"model_type": "compare", "features": ["age"]}})
    assert 'Replay request_config: {"features": ["age"], "model_type": "compare"}' in notes


# ── Export writers: headers, LaTeX rows, Markdown cells, journal numbers ──


def _cohort_with_multiline_groups() -> str:
    rng = np.random.default_rng(0)
    n = 120
    histology = np.where(rng.random(n) < 0.5, "Adenocarcinoma\n(NOS)", "Squamous")
    lines = ["os_months,os_event,age,histology"]
    for index in range(n):
        lines.append(f'{rng.exponential(30) + 0.5:.2f},{int(rng.random() < 0.6)},{rng.normal(60, 10):.1f},"{histology[index]}"')
    upload = client.post("/api/upload", files={"file": ("cohort.csv", "\n".join(lines).encode(), "text/csv")})
    assert upload.status_code == 200, upload.text
    return upload.json()["dataset_id"]


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "markdown", "latex", "docx"])
def test_server_produced_column_names_with_line_breaks_export_with_spaces(fmt: str) -> None:
    dataset_id = _cohort_with_multiline_groups()
    table = client.post("/api/cohort-table", json={"dataset_id": dataset_id, "variables": ["age"], "group_column": "histology"})
    assert table.status_code == 200, table.text
    analysis = table.json()["analysis"]
    assert any("\n" in column for column in analysis["columns"])

    response = client.post(
        "/api/export-table",
        json={"rows": analysis["rows"], "columns": analysis["columns"], "format": fmt, "style": "plain", "caption": "Cohort summary"},
    )
    assert response.status_code == 200, response.text
    if fmt == "xlsx":
        from openpyxl import load_workbook

        worksheet = load_workbook(io.BytesIO(response.content)).active
        text = "\n".join(str(cell.value) for row in worksheet.iter_rows() for cell in row if cell.value is not None)
    elif fmt == "docx":
        text = zipfile.ZipFile(io.BytesIO(response.content)).read("word/document.xml").decode("utf-8")
    else:
        text = response.text
    assert "Adenocarcinoma (NOS)" in text
    assert "Adenocarcinoma\n(NOS)" not in text


def test_latex_rows_cannot_be_read_as_optional_arguments() -> None:
    hazard = "Hazard\nratio\t(HR)"
    response = client.post(
        "/api/export-table",
        json={
            "rows": [{"CI": "[0.61, 0.74]", hazard: 1.2}, {"CI": "*0.50 to 0.90", hazard: 1.1}],
            "columns": ["CI", hazard],
            "format": "latex",
            "style": "plain",
        },
    )
    assert response.status_code == 200, response.text
    lines = response.text.splitlines()
    start = lines.index("\\toprule")
    end = lines.index("\\bottomrule")
    table_rows = [line for line in lines[start + 1 : end] if line != "\\midrule"]
    # "\\" and "\midrule" read a following "[" (or "*") as their own argument; an empty group ends that.
    assert table_rows and all(line.startswith("{}") for line in table_rows), table_rows
    assert table_rows[0] == "{}CI & Hazard ratio (HR) \\\\"


def _gfm_cells(row: str) -> list[str]:
    """Split a Markdown table row the way a renderer that honours every backslash escape does."""
    cells, current, escaped = [], "", False
    for character in row.strip().strip("|"):
        if escaped:
            current += character
            escaped = False
        elif character == "\\":
            escaped = True
        elif character == "|":
            cells.append(current.strip())
            current = ""
        else:
            current += character
    cells.append(current.strip())
    return cells


def test_markdown_cells_keep_pipes_backslashes_and_markup_literal() -> None:
    rows = [{"Term": "a\\|b", "Level": "x|y", "Note": "<b>bold</b> & p<0.05", "C:\\path": 1}]
    response = client.post("/api/export-table", json={"rows": rows, "format": "markdown", "style": "plain"})
    assert response.status_code == 200, response.text
    table = [line for line in response.text.splitlines() if line.startswith("|")]
    header, body = _gfm_cells(table[0]), _gfm_cells(table[2])
    assert header == ["Term", "Level", "Note", "C:\\path"]
    assert body == ["a\\|b", "x|y", "&lt;b&gt;bold&lt;/b&gt; &amp; p&lt;0.05", "1"]


def test_journal_style_formats_whole_numbers_of_measurement_columns() -> None:
    # The browser sends 1.0 as 1 and 0.0 as 0; counts and ranks stay whole numbers.
    rows = [
        {"Rank": 1, "Model": "A", "C-index": 0.7123, "P value": 0.0312, "Events, n": 45, "Training Time, ms": 1234.5},
        {"Rank": 2, "Model": "B", "C-index": 1, "P value": 1, "Events, n": 40, "Training Time, ms": 1234},
        {"Rank": 3, "Model": "C", "C-index": 0.5, "P value": 0, "Events, n": 12, "Training Time, ms": 99.0},
    ]
    csv_text = client.post("/api/export-table", json={"rows": rows, "format": "csv", "style": "journal"}).text
    import csv as csv_module

    parsed = list(csv_module.reader(io.StringIO(csv_text.lstrip("\ufeff"))))
    header_index = next(index for index, row in enumerate(parsed) if row and row[0] == "Rank")
    body = parsed[header_index + 1 :]
    assert body[1] == ["2", "B", "1.000", "1.000", "40", "1234.000"]
    assert body[2] == ["3", "C", "0.500", "<0.001", "12", "99.000"]

    from openpyxl import load_workbook

    xlsx = client.post("/api/export-table", json={"rows": rows, "format": "xlsx", "style": "journal"})
    worksheet = load_workbook(io.BytesIO(xlsx.content)).active
    second = [cell for cell in worksheet[3]]
    assert second[0].value == 2 and second[0].number_format == "General"
    assert second[2].value == 1.0 and second[2].number_format == "0.000"
    assert second[3].value == 1.0 and second[3].number_format == "0.000"
    assert second[4].value == 40 and second[4].number_format == "General"
    third = [cell for cell in worksheet[4]]
    assert third[3].value == "<0.001"


def test_journal_p_values_just_below_the_threshold_never_print_as_it() -> None:
    for value in (0.0499999951, 0.04999999999999999):
        text = app_module._format_journal_p_value(value)
        assert float(text) < 0.05, text


# ── Header clean-up keeps the names that needed none ────────────


def test_cleaned_headers_never_take_the_name_of_an_untouched_column() -> None:
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["Overall survival\n(months)", "Overall survival (months)", "Death", "Age\nat\ndiagnosis"])
    for index in range(60):
        worksheet.append([1000.0 + index, float(index + 1), index % 2, 50.0 + index])
    buffer = io.BytesIO()
    workbook.save(buffer)

    upload = client.post("/api/upload", files={"file": ("headers.xlsx", buffer.getvalue(), "application/octet-stream")})
    assert upload.status_code == 200, upload.text
    names = [column["name"] for column in upload.json()["columns"]]
    assert names == ["Overall survival (months)_2", "Overall survival (months)", "Death", "Age at diagnosis"]
    frame = app_module.store.get(upload.json()["dataset_id"], copy_dataframe=False).dataframe
    assert frame["Overall survival (months)"].tolist()[:3] == [1.0, 2.0, 3.0]
    assert frame["Overall survival (months)_2"].tolist()[:3] == [1000.0, 1001.0, 1002.0]
